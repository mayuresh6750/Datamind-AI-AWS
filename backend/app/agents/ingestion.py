"""Ingestion Agent — format detection, parsing, column profiling, merge plan."""
import csv
import io
import json
import re
from typing import Any

import pyarrow as pa
import pyarrow.parquet as pq
import structlog
from openpyxl import load_workbook

from app.core.llm import LLMClient
from app.models.schemas import (
    ColumnClassificationResponse,
    ColumnProfile,
    MergePlan,
)

log = structlog.get_logger()

# ─── Classification thresholds ────────────────────────────────────────────────
MAX_CARDINALITY_FOR_STRUCTURED = 50
MIN_AVG_LENGTH_FOR_TEXTUAL = 40
LLM_CONFIDENCE_THRESHOLD = 0.7  # columns below this get sent to LLM

# Date patterns (common formats)
_DATE_PATTERNS = [
    re.compile(r"^\d{4}[-/]\d{1,2}[-/]\d{1,2}"),                   # YYYY-MM-DD or YYYY/MM/DD
    re.compile(r"^\d{1,2}[-/]\d{1,2}[-/]\d{2,4}"),                 # DD/MM/YYYY or MM-DD-YYYY
    re.compile(r"^\d{4}[-/]\d{1,2}[-/]\d{1,2}[T ]\d{1,2}:\d{2}"), # ISO timestamp
    re.compile(r"^\w{3}\s+\d{1,2},?\s+\d{4}"),                     # Jan 15, 2024
]

# Column name hints for semantic type detection
_DATE_NAME_HINTS = {"date", "time", "created", "updated", "timestamp", "datetime", "dob", "born", "expired", "expiry"}
_ID_NAME_HINTS = {"id", "uuid", "key", "code", "number", "no", "num"}

# ─── LLM classification prompt ───────────────────────────────────────────────
_CLASSIFY_PROMPT = """\
Classify each column into exactly ONE of three classes:
- "numeric": values that represent quantities, amounts, counts, scores, or measurements that can be meaningfully aggregated (SUM, AVG, MAX, MIN). e.g. price, revenue, quantity, age, score.
- "structured": names, labels, categories, identifiers, dates/times, codes, or any non-aggregatable values that are still useful for GROUP BY or filtering in SQL — even if they look like numbers (e.g. ZIP codes, phone numbers, IDs, product categories, regions, dates).
- "textual": free-form, unstructured text with high cardinality and long strings — only for things like product descriptions, comments, reviews, notes. These get vectorized for semantic search.

Columns to classify:
{columns_block}

Return ONLY valid JSON matching this schema:
{schema}
""".strip()


class IngestionAgent:
    """Detects format, parses data, profiles columns, proposes merge plan."""

    def __init__(self, llm: LLMClient | None = None) -> None:
        self.llm = llm or LLMClient()

    # ── Format detection ──────────────────────────────────────────────────────

    def detect_format(self, filename: str, content: bytes) -> str:
        ext = filename.rsplit(".", 1)[-1].lower()
        if ext in ("csv", "xlsx", "json", "parquet"):
            return ext
        # fallback: sniff CSV
        try:
            sample = content[:2048].decode("utf-8", errors="replace")
            csv.Sniffer().sniff(sample)
            return "csv"
        except Exception:
            return "unknown"

    # ── Parsers ───────────────────────────────────────────────────────────────

    def parse_csv(self, content: bytes) -> list[dict]:
        text = content.decode("utf-8-sig", errors="replace")
        reader = csv.DictReader(io.StringIO(text, newline=""))
        # csv.DictReader stores overflow values under key None; strip those artifacts
        return [{k: v for k, v in row.items() if k is not None} for row in reader]

    def parse_json(self, content: bytes) -> list[dict]:
        data = json.loads(content)
        if isinstance(data, list):
            return data
        return [data]

    def parse_xlsx(self, content: bytes) -> list[dict]:
        workbook = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
        rows: list[dict] = []

        for sheet in workbook.worksheets:
            values = list(sheet.iter_rows(values_only=True))
            if not values:
                continue

            headers = [str(value).strip() if value is not None else "" for value in values[0]]
            if not any(headers):
                continue

            for raw_row in values[1:]:
                if raw_row is None or not any(value is not None and value != "" for value in raw_row):
                    continue

                row = {
                    header: value
                    for header, value in zip(headers, raw_row, strict=False)
                    if header
                }
                if row:
                    rows.append(row)

        return rows

    def parse_parquet(self, content: bytes) -> list[dict]:
        table = pq.read_table(pa.BufferReader(content))
        return table.to_pylist()

    def parse_content(self, fmt: str, content: bytes) -> list[dict]:
        if fmt == "csv":
            return self.parse_csv(content)
        if fmt == "json":
            return self.parse_json(content)
        if fmt == "xlsx":
            return self.parse_xlsx(content)
        if fmt == "parquet":
            return self.parse_parquet(content)
        raise ValueError(f"Unsupported file format: {fmt}")

    # ── Column classification helpers ─────────────────────────────────────────

    @staticmethod
    def _is_date_value(value: str) -> bool:
        """Check if a string looks like a date/timestamp."""
        return any(pattern.match(value) for pattern in _DATE_PATTERNS)

    @staticmethod
    def _col_name_tokens(col_name: str) -> set[str]:
        """Split column name into lowercase tokens for hint matching."""
        return set(col_name.lower().replace("-", "_").split("_"))

    @staticmethod
    def _classify_column_heuristic(
        col_name: str,
        non_null_values: list,
        total_rows: int,
    ) -> tuple[str, str, float]:
        """Classify a single column using rule-based heuristics.

        Returns (classification, dtype, confidence).
        """
        if not non_null_values:
            return ("ambiguous", "text", 0.3)

        str_values = [str(v) for v in non_null_values]
        name_tokens = IngestionAgent._col_name_tokens(col_name)
        distinct_count = len(set(str_values))
        cardinality_ratio = distinct_count / total_rows if total_rows > 0 else 0

        # ── 1. Numeric check ──────────────────────────────────────────────────
        numeric_count = 0
        total_str_len = 0
        for v in str_values:
            try:
                float(v.replace(",", ""))
                numeric_count += 1
            except ValueError:
                total_str_len += len(v)

        numeric_ratio = numeric_count / len(str_values)

        if numeric_ratio >= 0.9:
            # High numeric ratio — but is it an ID or a real measure?
            has_id_hint = bool(name_tokens & _ID_NAME_HINTS)

            if has_id_hint and cardinality_ratio > 0.5:
                # Numeric values + ID-like name + high cardinality → identifier, not a measure
                return ("structured", "text", 0.85)

            # Check if values are purely integer-like
            int_count = sum(1 for v in str_values if v.replace(",", "").replace("-", "").replace("+", "").isdigit())
            if int_count / len(str_values) >= 0.9:
                return ("numeric", "integer", 0.95)

            return ("numeric", "float", 0.95)

        # ── 2. Date / temporal check ──────────────────────────────────────────
        has_date_name = bool(name_tokens & _DATE_NAME_HINTS)
        sample_for_date = str_values[:20]
        date_count = sum(1 for v in sample_for_date if IngestionAgent._is_date_value(v.strip()))
        date_ratio = date_count / len(sample_for_date) if sample_for_date else 0

        if date_ratio >= 0.7 or (has_date_name and date_ratio >= 0.3):
            return ("structured", "datetime", 0.90)

        # ── 3. Boolean check ─────────────────────────────────────────────────
        bool_vals = {"true", "false", "yes", "no", "0", "1", "t", "f", "y", "n"}
        unique_lower = set(v.lower().strip() for v in str_values)
        if unique_lower and unique_lower.issubset(bool_vals) and len(unique_lower) <= 3:
            return ("numeric", "boolean", 0.95)

        # ── 4. Categorical/structured text (low cardinality, non-numeric) ────────────────
        if cardinality_ratio < 0.05 or distinct_count <= 20:
            return ("structured", "categorical", 0.85)

        # ── 5. Free-form text (high avg length) ──────────────────────────────
        non_numeric_count = len(str_values) - numeric_count
        avg_length = total_str_len / non_numeric_count if non_numeric_count > 0 else 0.0
        if avg_length >= MIN_AVG_LENGTH_FOR_TEXTUAL:
            return ("textual", "text", 0.90)

        # ── 6. Moderate cardinality non-numeric ──────────────────────────────
        if distinct_count <= MAX_CARDINALITY_FOR_STRUCTURED:
            return ("structured", "categorical", 0.70)

        # ── 7. Fallback ─────────────────────────────────────────────────────
        return ("structured", "text", 0.4)

    async def _classify_columns_llm(
        self,
        ambiguous_columns: list[dict],
    ) -> dict[str, tuple[str, str]]:
        """Call LLM to classify columns where heuristics were uncertain.

        Args:
            ambiguous_columns: list of dicts with keys 'name' and 'samples'.

        Returns:
            dict mapping column name → (classification, dtype).
            Empty dict on any failure.
        """
        if not ambiguous_columns:
            return {}

        # Build the columns block for the prompt
        lines = []
        for col in ambiguous_columns:
            samples_str = ", ".join(repr(s) for s in col["samples"][:5])
            lines.append(f"- {col['name']}: [{samples_str}]")

        prompt = _CLASSIFY_PROMPT.format(
            columns_block="\n".join(lines),
            schema=ColumnClassificationResponse.model_json_schema(),
        )

        try:
            result = await self.llm.structured_completion(
                prompt,
                ColumnClassificationResponse,
                agent_name="ingestion_classify",
            )
            return {
                item.name: (item.classification, item.dtype)
                for item in result.columns
                if item.classification in ("numeric", "structured", "textual")
            }
        except Exception as exc:
            log.warning("column_classification_llm_failed", error=str(exc))
            return {}

    # ── Column profiling (hybrid: heuristic + LLM) ───────────────────────────

    async def profile_columns(self, rows: list[dict]) -> list[ColumnProfile]:
        if not rows:
            return []

        columns = list(rows[0].keys())
        heuristic_results: list[dict] = []

        # Phase 1: Heuristic classification for every column
        for col in columns:
            values = [r.get(col) for r in rows]
            non_null = [v for v in values if v is not None and v != ""]
            null_count = len(values) - len(non_null)
            str_non_null = [str(v) for v in non_null]
            distinct_count = len(set(str_non_null))

            classification, dtype, confidence = self._classify_column_heuristic(
                col, non_null, len(rows),
            )

            # Compute avg_length for non-numeric string values
            total_str_len = 0
            numeric_count = 0
            for v in non_null:
                try:
                    float(str(v).replace(",", ""))
                    numeric_count += 1
                except ValueError:
                    total_str_len += len(str(v))
            non_numeric_count = len(non_null) - numeric_count
            avg_length = total_str_len / non_numeric_count if non_numeric_count > 0 else 0.0

            heuristic_results.append({
                "name": col,
                "dtype": dtype,
                "null_count": null_count,
                "distinct_count": distinct_count,
                "avg_length": avg_length if avg_length > 0 else None,
                "classification": classification,
                "confidence": confidence,
                "samples": [str(v) for v in non_null[:5]],
            })

        # Phase 2: LLM refinement for low-confidence columns
        ambiguous = [
            {"name": r["name"], "samples": r["samples"]}
            for r in heuristic_results
            if r["confidence"] < LLM_CONFIDENCE_THRESHOLD
        ]

        llm_overrides: dict[str, tuple[str, str]] = {}
        if ambiguous:
            log.info(
                "column_classification_calling_llm",
                ambiguous_count=len(ambiguous),
                columns=[c["name"] for c in ambiguous],
            )
            llm_overrides = await self._classify_columns_llm(ambiguous)

        # Build final profiles, merging LLM results where available
        profiles: list[ColumnProfile] = []
        for r in heuristic_results:
            classification = r["classification"]
            dtype = r["dtype"]

            if r["name"] in llm_overrides:
                llm_class, llm_dtype = llm_overrides[r["name"]]
                classification = llm_class
                dtype = llm_dtype
                log.info(
                    "column_classification_llm_override",
                    column=r["name"],
                    heuristic=r["classification"],
                    llm=classification,
                )

            profiles.append(ColumnProfile(
                name=r["name"],
                dtype=dtype,
                null_count=r["null_count"],
                distinct_count=r["distinct_count"],
                avg_length=r["avg_length"],
                classification=classification,
            ))

        return profiles

    # ── Merge plan ────────────────────────────────────────────────────────────

    def propose_merge_plan(
        self,
        file_profiles: list[tuple[str, list[ColumnProfile]]],
        all_rows: list[dict] | None = None,
    ) -> MergePlan:
        if len(file_profiles) == 1:
            return MergePlan(
                strategy="single",
                confidence=0.99,
                column_profiles=file_profiles[0][1],
            )

        # Check if all files share the same columns (UNION candidate)
        all_col_sets = [set(p.name for p in profiles) for _, profiles in file_profiles]
        if all(s == all_col_sets[0] for s in all_col_sets):
            # For union, use the profiles from the first file (already profiled).
            # We'll re-profile the combined set in the pipeline if needed.
            all_profiles = file_profiles[0][1]
            return MergePlan(strategy="union", confidence=0.92, column_profiles=all_profiles)

        # Different schemas — JOIN candidate; find common key column
        common_cols = all_col_sets[0].intersection(*all_col_sets[1:])
        join_key = next(iter(common_cols), None)
        all_profiles = [p for _, profiles in file_profiles for p in profiles]
        return MergePlan(
            strategy="join" if join_key else "separate",
            join_key=join_key,
            confidence=0.75 if join_key else 0.55,
            column_profiles=all_profiles,
        )

    def to_json_records(self, rows: list[dict]) -> list[dict]:
        return rows

